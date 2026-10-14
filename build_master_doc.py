#!/usr/bin/env python3
"""
build_master_doc.py — Sara's Snowflake note builder (recreated Class 31).

What it does
------------
1. Converts a single class's markdown notes file into a styled .docx
   (`Snowflake-ClassNN-Notes.docx`), matching the house style.
2. Rebuilds the running master doc `Snowflake-Classes-21-NN-Master-Notes.docx`
   by concatenating every per-class markdown file it can find (class21..classNN)
   plus a title page and an auto-generated Contents list.

Markdown supported (intentionally small, matches how the notes are written):
  # H1            -> Heading 1   (class title)
  ## H2           -> Heading 2
  ### H3          -> Heading 3
  #### H4         -> Heading 4
  > quote         -> Quote style (callouts / honest-labeling boxes)
  - bullet        -> List Bullet   (|  -> nested, two spaces indent = List Bullet 2)
  1. numbered     -> List Number
  | a | b |       -> table (markdown pipe table; a --- separator row is skipped)
  ```code```      -> monospace block (Consolas, light grey shading)
  plain line      -> Normal
Inline:  **bold**, *italic*, `code`  are rendered as runs.

Usage
-----
  python build_master_doc.py --class-md class31/snowflake-class31-notes.md --n 31
  python build_master_doc.py --rebuild-only --n 31      # just rebuild master
"""

import argparse
import glob
import os
import re
from docx import Document
from docx.shared import Pt, RGBColor
from docx.enum.text import WD_ALIGN_PARAGRAPH
from docx.oxml.ns import qn
from docx.oxml import OxmlElement

NAVY = RGBColor(0x1F, 0x38, 0x64)
CODE_GREY = "F2F2F2"
INLINE_RE = re.compile(r"(\*\*.+?\*\*|\*.+?\*|`.+?`)")


# ----------------------------- inline formatting -----------------------------
def add_inline(paragraph, text):
    """Render **bold**, *italic*, `code` as runs inside a paragraph."""
    for part in INLINE_RE.split(text):
        if not part:
            continue
        if part.startswith("**") and part.endswith("**"):
            r = paragraph.add_run(part[2:-2]); r.bold = True
        elif part.startswith("*") and part.endswith("*"):
            r = paragraph.add_run(part[1:-1]); r.italic = True
        elif part.startswith("`") and part.endswith("`"):
            r = paragraph.add_run(part[1:-1]); r.font.name = "Consolas"
        else:
            paragraph.add_run(part)


def shade(paragraph, hex_fill):
    pPr = paragraph._p.get_or_add_pPr()
    shd = OxmlElement("w:shd")
    shd.set(qn("w:val"), "clear")
    shd.set(qn("w:fill"), hex_fill)
    pPr.append(shd)


# ----------------------------- table rendering -------------------------------
def add_table(doc, rows):
    rows = [r for r in rows if not re.match(r"^\s*\|?[\s:-]+\|", r) or "---" not in r]
    cells = [[c.strip() for c in r.strip().strip("|").split("|")] for r in rows]
    cells = [c for c in cells if not all(set(x) <= set("-: ") for x in c)]
    if not cells:
        return
    ncols = max(len(r) for r in cells)
    t = doc.add_table(rows=0, cols=ncols)
    t.style = "Light Grid Accent 1"
    for ri, row in enumerate(cells):
        tr = t.add_row().cells
        for ci in range(ncols):
            txt = row[ci] if ci < len(row) else ""
            p = tr[ci].paragraphs[0]
            add_inline(p, txt)
            if ri == 0:
                for run in p.runs:
                    run.bold = True


# ----------------------------- markdown -> docx ------------------------------
def render_markdown(doc, md_text, skip_frontmatter_byline=True):
    """Append rendered markdown blocks to `doc`. Returns the H1 title found."""
    lines = md_text.splitlines()
    i = 0
    h1_title = None
    in_code = False
    code_buf = []
    table_buf = []

    def flush_table():
        if table_buf:
            add_table(doc, table_buf[:])
            table_buf.clear()

    while i < len(lines):
        line = lines[i]
        stripped = line.strip()

        # fenced code blocks
        if stripped.startswith("```"):
            if in_code:
                p = doc.add_paragraph()
                shade(p, CODE_GREY)
                r = p.add_run("\n".join(code_buf))
                r.font.name = "Consolas"; r.font.size = Pt(9)
                code_buf = []; in_code = False
            else:
                flush_table(); in_code = True
            i += 1; continue
        if in_code:
            code_buf.append(line); i += 1; continue

        # tables
        if stripped.startswith("|"):
            table_buf.append(stripped); i += 1; continue
        else:
            flush_table()

        if not stripped:
            i += 1; continue

        # headings
        if stripped.startswith("#"):
            level = len(stripped) - len(stripped.lstrip("#"))
            text = stripped[level:].strip()
            if level == 1:
                h1_title = text
                p = doc.add_paragraph(style="Heading 1"); add_inline(p, text)
            else:
                lvl = min(level, 4)
                p = doc.add_paragraph(style=f"Heading {lvl}"); add_inline(p, text)
            i += 1; continue

        # blockquote (callout) — gather consecutive > lines
        if stripped.startswith(">"):
            q_lines = []
            while i < len(lines) and lines[i].strip().startswith(">"):
                q_lines.append(lines[i].strip().lstrip(">").strip())
                i += 1
            p = doc.add_paragraph(style="Quote")
            add_inline(p, " ".join(l for l in q_lines if l))
            continue

        # bullets (support one nesting level via leading spaces)
        m = re.match(r"^(\s*)[-*]\s+(.*)$", line)
        if m:
            indent = len(m.group(1))
            style = "List Bullet 2" if indent >= 2 else "List Bullet"
            p = doc.add_paragraph(style=style); add_inline(p, m.group(2))
            i += 1; continue

        # numbered
        m = re.match(r"^\s*\d+\.\s+(.*)$", line)
        if m:
            p = doc.add_paragraph(style="List Number"); add_inline(p, m.group(1))
            i += 1; continue

        # horizontal rule -> skip (visual separator)
        if set(stripped) <= set("-") and len(stripped) >= 3:
            i += 1; continue

        # byline line in per-class frontmatter — skip when embedding in master
        if skip_frontmatter_byline and stripped.startswith("*Sara —"):
            i += 1; continue

        # plain paragraph
        p = doc.add_paragraph(); add_inline(p, stripped)
        i += 1

    flush_table()
    return h1_title


# ----------------------------- doc scaffolding -------------------------------
def title_page(doc, last_n):
    def centered(text, size, bold=False, color=None):
        p = doc.add_paragraph(); p.alignment = WD_ALIGN_PARAGRAPH.CENTER
        r = p.add_run(text); r.bold = bold; r.font.size = Pt(size)
        if color: r.font.color.rgb = color
        return p
    centered("Snowflake Training — Master Notes", 30, bold=True, color=NAVY)
    centered(f"Classes 21 – {last_n}", 18, color=NAVY)
    centered("Sara — SnowPro Core preparation", 12)
    doc.add_paragraph()


def build_per_class_docx(md_path, n, out_dir):
    doc = Document()
    with open(md_path) as f:
        render_markdown(doc, f.read(), skip_frontmatter_byline=False)
    out = os.path.join(out_dir, f"Snowflake-Class{n}-Notes.docx")
    doc.save(out)
    return out


def rebuild_master(repo_dir, last_n):
    """Concatenate all class markdown files 21..last_n into the master doc."""
    doc = Document()
    title_page(doc, last_n)

    # discover per-class markdown files
    found = {}
    for md in glob.glob(os.path.join(repo_dir, "**", "snowflake-class*-notes.md"), recursive=True):
        m = re.search(r"class(\d+)-notes\.md$", md)
        if m:
            found[int(m.group(1))] = md
    classes = sorted(k for k in found if 21 <= k <= last_n)

    # Contents
    doc.add_paragraph("Contents", style="Heading 1")
    titles = {}
    for n in classes:
        with open(found[n]) as f:
            first = next((l for l in f if l.startswith("# ")), f"# Class {n}")
        titles[n] = first[2:].strip()
        doc.add_paragraph(titles[n], style="List Number")

    # Body
    for n in classes:
        doc.add_page_break()
        with open(found[n]) as f:
            render_markdown(doc, f.read(), skip_frontmatter_byline=True)

    out = os.path.join(repo_dir, f"Snowflake-Classes-21-{last_n}-Master-Notes.docx")
    doc.save(out)
    return out, classes


# ----------------------------------- main ------------------------------------
if __name__ == "__main__":
    ap = argparse.ArgumentParser()
    ap.add_argument("--class-md", help="path to this class's markdown notes")
    ap.add_argument("--n", type=int, required=True, help="latest class number")
    ap.add_argument("--repo-dir", default=os.path.dirname(os.path.abspath(__file__)))
    ap.add_argument("--rebuild-only", action="store_true")
    args = ap.parse_args()

    if args.class_md and not args.rebuild_only:
        per = build_per_class_docx(args.class_md, args.n, os.path.dirname(args.class_md) or ".")
        print(f"[per-class] wrote {per}")

    master, classes = rebuild_master(args.repo_dir, args.n)
    print(f"[master] wrote {master}")
    print(f"[master] included classes: {classes}")
